"""Moliyaviy hisobotlar — ekran va eksport (CSV/Excel) bitta funksiyadan.

Summalar faqat tasdiqlangan operatsiyalardan; qoralama/bekor qilinganlar
kirmaydi. Valyutalar hech qachon bir summaga qo'shilmaydi — jami qiymatlar
`by_currency` lug'atida alohida.
"""
from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from typing import Optional

from fastapi import HTTPException
from sqlalchemy import or_
from sqlalchemy.orm import Session, joinedload

from .. import models, tz
from . import calc

ZERO = Decimal("0")
FT = models.FinanceTransaction

METHOD_LABELS = {"cash": "Naqd", "card": "Karta", "bank_transfer": "Bank o'tkazmasi",
                 "online": "Onlayn (Payme/Click/Uzum)", "unknown": "Belgilanmagan", None: "—"}
KIND_LABELS = {"income": "Kirim", "expense": "Xarajat", "refund": "Qaytarim",
               "adjustment": "Tuzatish", "transfer": "O'tkazma"}
STATUS_LABELS = {"draft": "Qoralama", "confirmed": "Tasdiqlangan", "cancelled": "Bekor qilingan"}
MONTHS = ["Yanvar", "Fevral", "Mart", "Aprel", "May", "Iyun", "Iyul", "Avgust",
          "Sentyabr", "Oktyabr", "Noyabr", "Dekabr"]

REPORTS = {
    "income":   "Kirimlar",
    "expenses": "Xarajatlar",
    "cashflow": "Pul oqimi",
    "accounts": "Kassa va hisoblar harakati",
    "debts":    "Mijozlar qarzdorligi",
    "overdue":  "Muddati o'tgan to'lovlar",
    "history":  "To'lov tarixi (mijoz/shartnoma)",
}


def _name(u):
    return (u.full_name or u.username) if u else None


def default_range(date_from: Optional[date], date_to: Optional[date]):
    today = tz.today()
    date_to = date_to or today
    date_from = date_from or date_to.replace(day=1)
    if date_from > date_to:
        raise HTTPException(400, "Sana oralig'i noto'g'ri: boshlanish tugashdan keyin")
    if (date_to - date_from).days > 3700:
        raise HTTPException(400, "Sana oralig'i juda katta (10 yildan ko'p)")
    return date_from, date_to


def base_query(db: Session, f: dict, *, flows=None, statuses=("confirmed",)):
    q = db.query(FT).options(
        joinedload(FT.account), joinedload(FT.to_account), joinedload(FT.student), joinedload(FT.group),
        joinedload(FT.category), joinedload(FT.created_by), joinedload(FT.responsible), joinedload(FT.origin),
    )
    if statuses:
        q = q.filter(FT.status.in_(statuses))
    if flows:
        q = q.filter(FT.flow.in_(flows))
    if f.get("date_from"):
        q = q.filter(FT.op_date >= f["date_from"])
    if f.get("date_to"):
        q = q.filter(FT.op_date <= f["date_to"])
    if f.get("method"):
        q = q.filter(FT.method == f["method"])
    if f.get("account_id"):
        q = q.filter(or_(FT.account_id == f["account_id"], FT.to_account_id == f["account_id"]))
    if f.get("category_id"):
        q = q.filter(FT.category_id == f["category_id"])
    if f.get("responsible_id"):
        q = q.filter(or_(FT.responsible_id == f["responsible_id"], FT.created_by_id == f["responsible_id"]))
    if f.get("group_id"):
        q = q.filter(FT.group_id == f["group_id"])
    if f.get("student_id"):
        q = q.filter(FT.student_id == f["student_id"])
    return q


def _add(d: dict, cur: str, key: str, v):
    d.setdefault(cur, defaultdict(lambda: ZERO))
    d[cur][key] += Decimal(str(v or 0))


def _fix(d: dict) -> dict:
    return {cur: dict(v) for cur, v in d.items()}


def tx_row(tx) -> dict:
    return {
        "date": tx.op_date, "number": tx.number, "kind": KIND_LABELS.get(tx.kind, tx.kind),
        "status": STATUS_LABELS.get(tx.status, tx.status),
        "student": tx.student.full_name if tx.student else None,
        "group": tx.group.name if tx.group else None,
        "period": f"{MONTHS[tx.period_month - 1]} {tx.period_year}" if tx.period_month and tx.period_year else None,
        "category": tx.category.name if tx.category else None,
        "counterparty": tx.counterparty,
        "method": METHOD_LABELS.get(tx.method, tx.method),
        "account": tx.account.name if tx.account else "Belgilanmagan (tarixiy)",
        "amount": Decimal(str(tx.signed_amount)) if tx.kind != "transfer" else Decimal(str(tx.amount)),
        "currency": tx.currency,
        "responsible": _name(tx.responsible) or _name(tx.created_by),
        "origin": tx.origin.number if tx.origin else None,
        "reason": tx.reason, "notes": tx.notes,
    }


TX_COLS = [
    ("date", "Sana", "date"), ("number", "Raqam", "text"), ("kind", "Turi", "text"),
    ("student", "Mijoz", "text"), ("group", "Guruh / shartnoma", "text"), ("period", "Davr", "text"),
    ("method", "To'lov usuli", "text"), ("account", "Kassa/hisob", "text"),
    ("amount", "Summa", "money"), ("currency", "Valyuta", "text"),
    ("responsible", "Mas'ul", "text"), ("origin", "Asl operatsiya", "text"), ("notes", "Izoh", "text"),
]


def report_income(db, f):
    rows, by_cur = [], {}
    for tx in base_query(db, f, flows=("income", "refund")).order_by(FT.op_date, FT.id):
        rows.append(tx_row(tx))
        key = "refunds" if tx.flow == "refund" else "income"
        _add(by_cur, tx.currency, key, abs(Decimal(str(tx.signed_amount))) if key == "refunds" else tx.signed_amount)
        _add(by_cur, tx.currency, "net", tx.signed_amount)
    return {"columns": TX_COLS, "rows": rows, "totals": _fix(by_cur)}


def report_expenses(db, f):
    rows, by_cur, by_cat = [], {}, {}
    cols = [c for c in TX_COLS if c[0] not in ("student", "group", "period")]
    cols.insert(3, ("category", "Kategoriya", "text"))
    cols.insert(4, ("counterparty", "Oluvchi", "text"))
    for tx in base_query(db, f, flows=("expense",)).order_by(FT.op_date, FT.id):
        r = tx_row(tx)
        r["amount"] = -Decimal(str(tx.signed_amount))
        rows.append(r)
        _add(by_cur, tx.currency, "expense", r["amount"])
        cat = (tx.category.name if tx.category else None) or (tx.origin.category.name if tx.origin and tx.origin.category else "Kategoriyasiz")
        _add(by_cat.setdefault(cat, {}), tx.currency, "amount", r["amount"])
    breakdown = sorted(
        ({"category": k, "by_currency": _fix(v)} for k, v in by_cat.items()),
        key=lambda x: -float(x["by_currency"].get("UZS", {}).get("amount", 0)),
    )
    return {"columns": cols, "rows": rows, "totals": _fix(by_cur), "breakdown": breakdown}


def report_cashflow(db, f):
    d_from, d_to = f["date_from"], f["date_to"]
    monthly = f.get("granularity") == "month" or (d_to - d_from).days > 92
    buckets = {}
    for tx in base_query(db, f, flows=("income", "refund", "expense")):
        key = (tx.op_date.year, tx.op_date.month) if monthly else tx.op_date
        b = buckets.setdefault((key, tx.currency), defaultdict(lambda: ZERO))
        s = Decimal(str(tx.signed_amount))
        if tx.flow == "income":
            b["income"] += s
        elif tx.flow == "refund":
            b["refunds"] += -s
        else:
            b["expense"] += -s
        b["net"] += s
    keys = []
    if monthly:
        y, m = d_from.year, d_from.month
        while (y, m) <= (d_to.year, d_to.month):
            keys.append((y, m))
            m, y = (1, y + 1) if m == 12 else (m + 1, y)
    else:
        d = d_from
        while d <= d_to:
            keys.append(d)
            d += timedelta(days=1)
    currencies = sorted({c for (_, c) in buckets} | {"UZS"})
    rows, by_cur = [], {}
    for cur in currencies:
        for k in keys:
            b = buckets.get((k, cur), {})
            label = f"{MONTHS[k[1] - 1]} {k[0]}" if monthly else k
            rows.append({"period": label, "currency": cur, "income": b.get("income", ZERO),
                         "refunds": b.get("refunds", ZERO), "expense": b.get("expense", ZERO),
                         "net": b.get("net", ZERO)})
            for kk in ("income", "refunds", "expense", "net"):
                _add(by_cur, cur, kk, b.get(kk, ZERO))
    cols = [("period", "Davr", "text"), ("currency", "Valyuta", "text"), ("income", "Tushum", "money"),
            ("refunds", "Qaytarimlar", "money"), ("expense", "Xarajatlar", "money"),
            ("net", "Sof pul oqimi", "money")]
    return {"columns": cols, "rows": rows, "totals": _fix(by_cur), "granularity": "month" if monthly else "day"}


def account_effect(tx, account_id) -> Decimal:
    if tx.kind == "transfer":
        if tx.account_id == account_id:
            return -Decimal(str(tx.amount))
        if tx.to_account_id == account_id:
            return Decimal(str(tx.amount))
        return ZERO
    return Decimal(str(tx.signed_amount)) if tx.account_id == account_id else ZERO


def report_accounts(db, f):
    d_from, d_to = f["date_from"], f["date_to"]
    accounts = db.query(models.FinanceAccount).order_by(models.FinanceAccount.id).all()
    if f.get("account_id"):
        accounts = [a for a in accounts if a.id == f["account_id"]]
    rows, lines = [], []
    for a in accounts:
        txs = (
            db.query(FT).options(joinedload(FT.student), joinedload(FT.category), joinedload(FT.account),
                                 joinedload(FT.to_account), joinedload(FT.created_by), joinedload(FT.group),
                                 joinedload(FT.responsible), joinedload(FT.origin))
            .filter(FT.status == "confirmed", or_(FT.account_id == a.id, FT.to_account_id == a.id),
                    FT.op_date <= d_to)
            .order_by(FT.op_date, FT.id).all()
        )
        opening = Decimal(str(a.opening_balance))
        inflow = outflow = ZERO
        running = None
        for tx in txs:
            eff = account_effect(tx, a.id)
            if tx.op_date < d_from:
                opening += eff
                continue
            if running is None:
                running = opening
            running += eff
            if eff >= 0:
                inflow += eff
            else:
                outflow += -eff
            if f.get("account_id"):
                r = tx_row(tx)
                r["amount"] = eff
                r["balance"] = running
                lines.append(r)
        rows.append({"account": a.name, "kind": a.kind, "currency": a.currency, "opening": opening,
                     "inflow": inflow, "outflow": outflow, "closing": opening + inflow - outflow,
                     "is_active": a.is_active})
    cols = [("account", "Kassa/hisob", "text"), ("currency", "Valyuta", "text"),
            ("opening", "Davr boshidagi qoldiq", "money"), ("inflow", "Kirim", "money"),
            ("outflow", "Chiqim", "money"), ("closing", "Davr oxiridagi qoldiq", "money")]
    by_cur = {}
    for r in rows:
        for k in ("opening", "inflow", "outflow", "closing"):
            _add(by_cur, r["currency"], k, r[k])
    out = {"columns": cols, "rows": rows, "totals": _fix(by_cur)}
    if f.get("account_id"):
        out["lines_columns"] = [("date", "Sana", "date"), ("number", "Raqam", "text"), ("kind", "Turi", "text"),
                                ("student", "Mijoz", "text"), ("category", "Kategoriya", "text"),
                                ("amount", "Harakat", "money"), ("balance", "Qoldiq", "money"),
                                ("notes", "Izoh", "text")]
        out["lines"] = lines
    return out


def debt_rows(db, f):
    rows = []
    due_day = calc.get_int_setting(db, "due_day")
    today = tz.today()
    for s in calc.students_for_debts(db, group_id=f.get("group_id"), search=f.get("search")):
        sch = calc.student_schedule(db, s, today=today, due_day=due_day)
        t = sch["totals"]
        if t["net_total"] <= 0 and t["paid_total"] <= 0 and not sch["installments"]:
            continue
        rows.append({"student_id": s.id, "student": s.full_name, "phone": s.phone1,
                     "groups": ", ".join(sorted({c["group_name"] for c in sch["contracts"] if not c["left_at"]})) or "—",
                     "is_archived": bool(s.is_archived),
                     **t, "_installments": sch["installments"]})
    return rows


DEBT_COLS = [("student", "Mijoz", "text"), ("phone", "Telefon", "text"), ("groups", "Guruh(lar)", "text"),
             ("contract_total", "Shartnoma summasi", "money"), ("discounts", "Chegirmalar", "money"),
             ("paid_total", "To'langan", "money"), ("debt", "Qolgan qarz", "money"),
             ("overdue_debt", "Muddati o'tgan", "money"), ("not_due_debt", "Muddati kelmagan", "money"),
             ("advance", "Avans", "money"), ("next_due_date", "Keyingi muddat", "date"),
             ("overdue_days", "Kechikish (kun)", "int")]


def report_debts(db, f):
    rows = debt_rows(db, f)
    status = f.get("status")
    if status == "debtor":
        rows = [r for r in rows if r["debt"] > 0]
    elif status == "overdue":
        rows = [r for r in rows if r["overdue_debt"] > 0]
    elif status == "advance":
        rows = [r for r in rows if r["advance"] > 0]
    totals = defaultdict(lambda: ZERO)
    for r in rows:
        for k in ("contract_total", "discounts", "paid_total", "debt", "overdue_debt", "not_due_debt", "advance"):
            totals[k] += r[k]
    for r in rows:
        r.pop("_installments", None)
    return {"columns": DEBT_COLS, "rows": rows, "totals": {"UZS": dict(totals)}}


def report_overdue(db, f):
    rows = []
    totals = defaultdict(lambda: ZERO)
    for r in debt_rows(db, f):
        for i in r["_installments"]:
            if i["status"] != "overdue":
                continue
            rows.append({"student_id": r["student_id"], "student": r["student"], "phone": r["phone"],
                         "group": i["group_name"], "period": f"{MONTHS[i['month'] - 1]} {i['year']}",
                         "due_date": i["due_date"], "overdue_days": i["overdue_days"],
                         "net": i["net"], "paid": i["paid"], "remaining": i["remaining"]})
            totals["remaining"] += i["remaining"]
    rows.sort(key=lambda x: -x["overdue_days"])
    cols = [("student", "Mijoz", "text"), ("phone", "Telefon", "text"), ("group", "Guruh", "text"),
            ("period", "Davr", "text"), ("due_date", "To'lov muddati", "date"),
            ("overdue_days", "Kechikish (kun)", "int"), ("net", "Hisoblangan", "money"),
            ("paid", "To'langan", "money"), ("remaining", "Qoldiq qarz", "money")]
    return {"columns": cols, "rows": rows, "totals": {"UZS": dict(totals)}}


def report_history(db, f):
    if not f.get("student_id") and not f.get("group_id"):
        raise HTTPException(400, "Mijoz yoki guruh (shartnoma)ni tanlang")
    rows, by_cur = [], {}
    q = base_query(db, f, flows=("income", "refund"), statuses=("draft", "confirmed", "cancelled"))
    for tx in q.order_by(FT.op_date, FT.id):
        r = tx_row(tx)
        rows.append(r)
        if tx.status == "confirmed":
            _add(by_cur, tx.currency, "net", tx.signed_amount)
    cols = TX_COLS[:3] + [("status", "Holat", "text")] + TX_COLS[3:]
    return {"columns": cols, "rows": rows, "totals": _fix(by_cur)}


BUILDERS = {
    "income": report_income, "expenses": report_expenses, "cashflow": report_cashflow,
    "accounts": report_accounts, "debts": report_debts, "overdue": report_overdue, "history": report_history,
}


def build(db: Session, kind: str, f: dict) -> dict:
    if kind not in BUILDERS:
        raise HTTPException(404, "Bunday hisobot yo'q")
    f = dict(f)
    f["date_from"], f["date_to"] = default_range(f.get("date_from"), f.get("date_to"))
    if kind in ("debts", "overdue"):
        f.pop("date_from"), f.pop("date_to")
    out = BUILDERS[kind](db, f)
    out["title"] = REPORTS[kind]
    out["columns"] = [{"key": k, "label": l, "type": t} for k, l, t in out["columns"]]
    if "lines_columns" in out:
        out["lines_columns"] = [{"key": k, "label": l, "type": t} for k, l, t in out["lines_columns"]]
    return out


# ── Eksport ──────────────────────────────────────────────────────────────────

def _cell(v, t):
    if v is None:
        return ""
    if t == "date" and hasattr(v, "strftime"):
        return v.strftime("%d.%m.%Y")
    return v


def to_csv(rep: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([rep["title"]])
    w.writerow([c["label"] for c in rep["columns"]])
    for r in rep["rows"]:
        w.writerow([_cell(r.get(c["key"]), c["type"]) for c in rep["columns"]])
    for cur, tot in (rep.get("totals") or {}).items():
        w.writerow([f"Jami ({cur})"] + [f"{k}: {v}" for k, v in tot.items()])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = rep["title"][:31]
    ws.append([rep["title"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([c["label"] for c in rep["columns"]])
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for r in rep["rows"]:
        row = []
        for c in rep["columns"]:
            v = r.get(c["key"])
            if c["type"] == "money" and v is not None:
                v = float(v)
            row.append(_cell(v, c["type"]))
        ws.append(row)
        for i, c in enumerate(rep["columns"], start=1):
            if c["type"] == "money":
                ws.cell(row=ws.max_row, column=i).number_format = "#,##0.00"
    ws.append([])
    for cur, tot in (rep.get("totals") or {}).items():
        ws.append([f"Jami ({cur})"] + [f"{k}: {float(v):,.2f}".replace(",", " ") for k, v in tot.items()])
    for i, c in enumerate(rep["columns"], start=1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = 22 if c["type"] == "text" else 16
    if rep.get("breakdown"):
        ws2 = wb.create_sheet("Kategoriyalar")
        ws2.append(["Kategoriya", "Valyuta", "Summa"])
        for b in rep["breakdown"]:
            for cur, v in b["by_currency"].items():
                ws2.append([b["category"], cur, float(v["amount"])])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
